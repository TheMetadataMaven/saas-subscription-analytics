"""
Build a single multi-sheet Excel workbook from the sample CSVs so the whole
CloudSync Pro model can be imported into Power BI in one step
(Get Data -> Excel workbook -> check all sheets -> Load).

Run:  py build_workbook.py
Output: ../CloudSyncPro_data.xlsx  (one sheet per table)
"""
import csv
import re
import datetime
from pathlib import Path
from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
SRC = HERE / "sample_data"
OUT = HERE.parent / "CloudSyncPro_data.xlsx"

# Sheet order (Excel sheet names max 31 chars; these are fine).
TABLES = [
    "fct_mrr",
    "rpt_unit_economics",
    "rpt_churn_risk",
    "dim_customers",
    "dim_subscriptions",
    "rpt_cohort_retention",
]

DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")   # full date only (not YYYY-MM keys)


def coerce(value: str):
    """Convert a CSV string cell into a typed Python value for Excel."""
    if value is None or value == "":
        return None
    low = value.lower()
    if low == "true":
        return True
    if low == "false":
        return False
    if DATE_RE.match(value):
        try:
            return datetime.date.fromisoformat(value)
        except ValueError:
            pass
    # integer (no decimal point, optional sign)
    if re.fullmatch(r"-?\d+", value):
        try:
            return int(value)
        except ValueError:
            pass
    # float
    if re.fullmatch(r"-?\d*\.\d+", value):
        try:
            return float(value)
        except ValueError:
            pass
    return value  # leave as text (keeps codes like "2024-01", "C0001")


def main():
    wb = Workbook()
    wb.remove(wb.active)  # drop default sheet
    for table in TABLES:
        path = SRC / f"{table}.csv"
        ws = wb.create_sheet(title=table[:31])
        with path.open(newline="", encoding="utf-8") as f:
            reader = csv.reader(f)
            header = next(reader)
            ws.append(header)
            for row in reader:
                ws.append([coerce(c) for c in row])
        # widen columns a little for readability
        for i, name in enumerate(header, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, len(name) + 2)
        ws.freeze_panes = "A2"
        print(f"  {table:24} {ws.max_row - 1:>3} rows, {len(header)} cols")
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
